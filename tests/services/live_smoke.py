"""手动小样本验收，显式 --live 才访问真实语音接口。"""
import argparse
import json
import sys
import tempfile
import time
from pathlib import Path
from docx import Document
from audiobook_generator.services.files import FileStore
from audiobook_generator.services.projects import Projects
from audiobook_generator.services.tasks import TaskManager
from audiobook_generator.services.models import Generation, TaskRequest


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--live', action='store_true')
    parser.add_argument('--provider', choices=['openai', 'edge', 'minimax', 'qwen'], default='openai')
    parser.add_argument('--report', required=True)
    parser.add_argument('--kind', choices=['generate', 'split'], default='generate')
    args = parser.parse_args()
    if not args.live:
        parser.error('真实接口验收需要显式传入 --live；会产生少量合成请求。')
    with tempfile.TemporaryDirectory(prefix='studio-live-') as directory:
        root = Path(directory)
        doc = Document()
        doc.add_heading('第一章', level=1)
        doc.add_paragraph('欢迎来到有声书工坊。这是一段用于验证界面迁移的短音频。')
        path = root / '迁移验收.docx'
        doc.save(path)
        files = FileStore(root)
        upload = files.register(path)
        manager = TaskManager(root, files, Projects(root, files))
        try:
            task = manager.start(TaskRequest(request_id='manual-live-smoke', kind=args.kind, generation=Generation(
                file_ids=[upload['id']], engine={'provider':args.provider}, worker_count=1)))
            deadline = time.time() + 90
            while time.time() < deadline:
                task = manager.get(task['id'])
                if task['status'] not in {'starting','running'}:
                    break
                time.sleep(.5)
            if task['status'] in {'starting','running'}:
                task = manager.stop(task['id'])
            audio = [p for p in root.rglob('*.mp3') if '.studio' not in p.parts]
            report = {'provider':args.provider,'status':task['status'],'error':task['error'],
                      'audio_files':[{'name':p.name,'bytes':p.stat().st_size} for p in audio],
                      'progress':task['progress'], 'kind':args.kind,
                      'project': {key: task['result']['project'][key] for key in ('id','revision','confirmed')} if task['result'].get('project') else None}
            Path(args.report).write_text(json.dumps(report, ensure_ascii=False, indent=2))
            print(json.dumps(report, ensure_ascii=False))
            return 0 if task['status']=='succeeded' and (audio or task['result'].get('project')) else 1
        finally:
            manager.close()

if __name__ == '__main__':
    sys.exit(main())
