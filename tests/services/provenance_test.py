import json
import tempfile
import unittest
import wave
from dataclasses import replace
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch

from bs4 import BeautifulSoup
from docx import Document
from fastapi.testclient import TestClient

from audiobook_generator.api.app import create_app
from audiobook_generator.book_parsers.doc_book_parser import DocBookParser
from audiobook_generator.book_parsers.source_map import TrackedText, epub_source
from audiobook_generator.config.general_config import GeneralConfig
from audiobook_generator.core.audiobook_generator import AudiobookGenerator
from audiobook_generator.core.audio_tags import AudioTags
from audiobook_generator.services.common import read_json
from audiobook_generator.services.provenance import write_sources, write_timeline, write_plans, finalize, sha256


class ProvenanceTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name).resolve()
        self.source = self.root / '原书.docx'
        doc = Document()
        doc.add_paragraph('  第一段[12]，你好。  ')
        doc.add_paragraph('')
        doc.add_paragraph('第二段，原书内容。')
        doc.save(self.source)
        self.cfg = GeneralConfig(None)
        self.cfg.input_file = str(self.source)
        self.cfg.output_folder = str(self.root / 'output')
        self.cfg.tts = 'openai'
        self.cfg.newline_mode = 'double'
        self.cfg.title_mode = 'auto'
        self.cfg.chapter_start, self.cfg.chapter_end = 1, -1
        self.cfg.remove_reference_numbers = True
        self.cfg.output_format = 'wav'
        self.parser = DocBookParser(self.cfg)
        self.chapters = self.parser.get_chapters('   ')

    def prepared(self):
        self.assertTrue(write_sources(self.parser, self.chapters, self.cfg))
        output = Path(self.cfg.output_folder)
        self.audio = output / 'chapter.wav'
        with wave.open(str(self.audio), 'wb') as stream:
            stream.setnchannels(1); stream.setsampwidth(2); stream.setframerate(16000)
            stream.writeframes(b'\0\0' * 16000)
        return output

    def span(self, **kwargs):
        return {'start_sample':0, 'end_sample':16000, 'source_span':{'chapter_index':1, 'start':0, 'end':len(self.chapters[0][1])}, **kwargs}

    def test_research_source_changed_after_preflight_does_not_write_manifest(self):
        self.cfg.research_render_snapshot = {'book_sha256': 'changed', 'chapters': {'1': {}}, 'source_digests': {}}
        with self.assertRaisesRegex(ValueError, '原文在生成前发生变化'):
            write_sources(self.parser, self.chapters, self.cfg)
        self.assertFalse(Path(self.cfg.output_folder, 'manifest.json').exists())

    def test_docx_keeps_original_paragraph_identity_after_cleanup(self):
        output = self.prepared()
        mapping = read_json(output / 'source_map.json')
        paragraphs = mapping['paragraphs']
        self.assertEqual([p['locator']['paragraph_index'] for p in paragraphs], [0, 2])
        self.assertEqual(paragraphs[0]['original_text'], '  第一段[12]，你好。  ')
        self.assertNotIn('[12]', paragraphs[0]['cleaned_text'])
        self.assertEqual(paragraphs[0]['locator']['paragraph_char_start'], 2)
        manifest = read_json(output / 'manifest.json')
        self.assertEqual(sha256(output / manifest['source_file']), sha256(self.source))

    def test_uploaded_docx_keeps_original_title_in_catalog_and_manifest(self):
        from audiobook_generator.services.books import configs_for, read_book
        from audiobook_generator.services.files import FileStore
        from audiobook_generator.services.models import BookRequest, Generation
        import io
        files = FileStore(self.root)
        file = files.upload('展示书名.docx', io.BytesIO(self.source.read_bytes()))
        _, chapters = read_book(BookRequest(file_id=file['id']), files)
        self.assertEqual(chapters[0][0], '展示书名')
        cfg = configs_for(Generation(file_ids=[file['id']]), files, self.root / 'upload_output')[0]
        parser = DocBookParser(cfg)
        write_sources(parser, parser.get_chapters('   '), cfg)
        self.assertEqual(read_json(Path(cfg.output_folder) / 'manifest.json')['title'], '展示书名')

    def test_epub_inline_nodes_preserve_exact_original_text_and_node_position(self):
        soup = BeautifulSoup('<html><body><p>你好<b>世界</b> <i>朋友</i>。</p><p>第二段。</p></body></html>', 'lxml-xml')
        raw, paragraphs = epub_source(soup, 'Text/one.xhtml')
        self.assertEqual(raw, '你好世界 朋友。\n第二段。')
        self.assertEqual(paragraphs[0]['text'], '你好世界 朋友。')
        self.assertEqual(paragraphs[0]['locator']['node_path'], '/html[1]/body[1]/p[1]')
        self.assertEqual(paragraphs[0]['locator']['href'], 'Text/one.xhtml')
        self.assertEqual(raw[paragraphs[0]['raw_start']:paragraphs[0]['raw_end']], paragraphs[0]['text'])

    def test_regex_replacements_and_deleted_paragraphs_keep_valid_coordinates(self):
        raw = '甲[12]\n乙\n丙'
        records = [{'raw_start':a, 'raw_end':b, 'text':raw[a:b], 'locator':{}} for a,b in [(0,5),(6,7),(8,9)]]
        tracked = TrackedText(raw).sub(r'\[\d+\]', '').sub('乙', '').sub('丙', '丙丙')
        mapping = tracked.paragraphs(records, 1)
        self.assertEqual(tracked.text, '甲\n\n丙丙')
        self.assertEqual(mapping[1]['cleaned_text'], '')
        self.assertEqual(mapping[2]['cleaned_text'], '丙丙')
        self.assertTrue(all(0 <= p['start'] <= p['end'] <= len(tracked.text) for p in mapping))

    def test_timeline_links_actual_samples_to_paragraphs_and_rejects_changed_audio(self):
        output = self.prepared()
        write_timeline(self.audio, 1, [self.span()], sample_rate=16000)
        finalize(output)
        timeline = read_json(output / 'timeline.json')['chapters'][0]
        self.assertEqual(timeline['spans'][0]['paragraph_ids'], ['ch1:p1','ch1:p2'])
        self.assertEqual(timeline['spans'][0]['end_seconds'], 1)
        self.audio.write_bytes(b'changed')
        finalize(output, status='failed')
        self.assertEqual(read_json(output / 'timeline.json')['chapters'], [])

    def test_rerun_never_claims_old_audio_completed_this_render(self):
        output = self.prepared()
        write_timeline(self.audio, 1, [self.span()], sample_rate=16000)
        finalize(output)
        self.assertEqual(read_json(output / 'manifest.json')['completed_chapters'], [1])
        write_sources(self.parser, self.chapters, self.cfg)
        finalize(output, status='failed')
        self.assertEqual(read_json(output / 'manifest.json')['completed_chapters'], [])

    def test_cli_validation_failure_finalizes_partial_manifest(self):
        self.cfg.chapter_start = 2
        provider = Mock()
        provider.get_break_string.return_value = '   '
        with patch('audiobook_generator.core.audiobook_generator.get_tts_provider', return_value=provider):
            self.assertFalse(AudiobookGenerator(self.cfg).run())
        manifest = read_json(Path(self.cfg.output_folder) / 'manifest.json')
        self.assertEqual(manifest['status'], 'failed')
        self.assertEqual(manifest['completed_chapters'], [])

    def test_missing_source_coordinates_have_only_chapter_precision(self):
        output = self.prepared()
        write_timeline(self.audio, 1, [{'start_sample':0,'end_sample':16000}], sample_rate=16000)
        finalize(output)
        self.assertEqual(read_json(output / 'timeline.json')['alignment'], 'chapter')
        self.assertEqual(read_json(output / 'timeline.json')['chapters'][0]['spans'][0]['precision'], 'chapter')
        with self.assertRaises(ValueError):
            write_timeline(self.audio, 1, [self.span(source_span={'start':-1,'end':20})], sample_rate=16000)

    def test_effective_profiles_keep_all_roles_sharing_a_voice(self):
        self.cfg.role_portable_voice_snapshots = {'hero': {'id': 'hero', 'version': 'a' * 64}}
        output = self.prepared()
        from audiobook_generator.role_script.models import ChapterSynthesisPlan, SynthesisUnit, VoiceProfile
        params = {'voice': 'Dean', 'speed': 1.1}
        units = tuple(SynthesisUnit(f'1:{role}', 1, '原书', i, 2, role, name, 'hello',
            VoiceProfile(role, name, params)) for i, (role, name) in enumerate(
                [('narrator', '旁白'), ('hero', '主角')], 1))
        write_plans(output, [ChapterSynthesisPlan(1, '原书', units, 'wav', 'openai')])
        profiles = list(read_json(output / 'voices.json')['planned_profiles'].values())
        self.assertEqual({p['role_id'] for p in profiles}, {'narrator', 'hero'})
        self.assertTrue(all(p['params'] == params for p in profiles))
        self.assertEqual(profiles[0]['version'], profiles[1]['version'])
        self.assertEqual(read_json(output / 'voices.json')['role_portable_voice_snapshots'],
            self.cfg.role_portable_voice_snapshots)

    def test_quality_casting_receives_script_and_keeps_locked_voices(self):
        from tests.audiobook_generator.role_script.v2_io_test import _script
        from audiobook_generator.role_script.models import ChapterSynthesisPlan
        script, _ = _script()
        self.cfg.role_script = script
        self.cfg.audio_core_casting = {'test': True}
        plan = ChapterSynthesisPlan(1, '原书', (), 'wav', 'openai')
        with patch('audiobook_generator.audio_core.casting.apply_casting', return_value=plan) as casting, \
                patch('audiobook_generator.audio_core.pipeline.run_chapter', return_value={}):
            AudiobookGenerator(self.cfg)._run_audio_core(plan, None, str(self.root / 'output.wav'), None)
        self.assertEqual(casting.call_args.kwargs['locked']['char_1'], script.roles['char_1'].provider_params['voice'])

    def test_library_api_serves_original_and_audio_but_marks_replaced_audio_stale(self):
        output = self.prepared()
        write_timeline(self.audio, 1, [self.span()], sample_rate=16000)
        finalize(output)
        with TestClient(create_app(self.root, acquire_lock=False)) as client:
            result = client.get('/api/v1/library/source?folder=output').json()
            self.assertTrue(result['available'])
            self.assertEqual(client.get('/api/v1/files/' + result['source_file']['id']).content, self.source.read_bytes())
            self.assertEqual(client.get('/api/v1/files/' + result['timeline']['chapters'][0]['audio_file']['id']).status_code, 200)
            self.audio.write_bytes(b'replaced')
            stale = client.get('/api/v1/library/source?folder=output').json()['timeline']['chapters'][0]
            self.assertIsNone(stale['audio_file'])
            self.assertTrue(stale['stale'])

    def test_plain_narration_uses_fixed_voice_and_exports_real_chunk_timeline(self):
        output = self.prepared()
        from audiobook_generator.role_script.audio_pipeline import build_audio_artifact
        calls = []
        def synthesize(text, profile, work, progress):
            calls.append((text, dict(profile.provider_params)))
            path = Path(work) / f'{len(calls)}.wav'
            with wave.open(str(path), 'wb') as stream:
                stream.setnchannels(1); stream.setsampwidth(2); stream.setframerate(16000)
                stream.writeframes(b'\0\0' * 1600)
            return build_audio_artifact(str(path), cache_key=str(len(calls)))
        text = self.chapters[0][1]
        self.cfg._portable_params = {'voice':'茉莉', 'speed':1.1}
        generator = AudiobookGenerator(self.cfg)
        generator._run_traceable_plain(1, '原书', text, SimpleNamespace(synthesize_segment=synthesize),
            str(self.audio), AudioTags('原书', '', '原书', 1))
        finalize(output)
        self.assertEqual(''.join(t for t,p in calls), text)
        self.assertTrue(all(p == self.cfg._portable_params for t,p in calls))
        timeline = read_json(output / 'timeline.json')['chapters'][0]
        self.assertEqual(len(timeline['spans']), len(calls))
        self.assertEqual(timeline['spans'][-1]['source_end'], len(text))
