"""工作台服务与 HTTP 回归；临时目录，不调用模型或真实 API。"""
import io
import json
import multiprocessing
import os
import signal
import subprocess
import sys
import tempfile
import time
import unittest
from pathlib import Path
from unittest.mock import patch

from fastapi.testclient import TestClient
from audiobook_generator.api.app import create_app
from audiobook_generator.services.books import configs_for
from audiobook_generator.services.common import ServiceError, RuntimeLock, atomic_json
from audiobook_generator.services.files import FileStore
from audiobook_generator.services.models import Generation, TaskRequest, ProjectEdit
from audiobook_generator.services.projects import Projects, encode
from audiobook_generator.services.settings import Settings, config_view, save_config
from audiobook_generator.services.tasks import TaskManager
from tests.audiobook_generator.role_script.v2_io_test import _script


def fake_worker(payload, result_path, log_path, ready):
    os.setsid()
    ready.set()
    Path(log_path).write_text('第一行\n第二行\n', encoding='utf-8')
    atomic_json(result_path, {'ok': True, 'result': {'offline': True}})


def tree_worker(payload, result_path, log_path, ready):
    os.setsid()
    ready.set()
    child = subprocess.Popen([sys.executable, '-c', 'import time; time.sleep(60)'])
    Path(log_path).write_text(str(child.pid))
    time.sleep(60)


def late_result_worker(payload, result_path, log_path, ready):
    os.setsid()
    def canceled(*_):
        atomic_json(result_path, {'ok': True, 'result': {'script': encode(_script()[0])}})
        raise SystemExit(0)
    signal.signal(signal.SIGTERM, canceled)
    ready.set()
    time.sleep(60)


class StudioTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.files = FileStore(self.root)
        self.file = self.files.upload('书.epub', io.BytesIO(b'book'))
        self.projects = Projects(self.root, self.files)

    def test_six_engine_mapping_and_no_prompt(self):
        for provider in ['openai', 'minimax', 'edge', 'qwen', 'chatterbox', 'voxcpm']:
            g = Generation(file_ids=[self.file['id']], engine={'provider': provider})
            cfg = configs_for(g, self.files, self.root)[0]
            self.assertEqual(cfg.tts, provider)
            self.assertTrue(cfg.no_prompt)
            self.assertEqual(cfg.worker_count, 1)
            self.assertEqual(cfg.newline_mode, 'double')
            if provider == 'edge':
                self.assertEqual(cfg.voice_rate, '+0%')
            if provider == 'chatterbox':
                self.assertEqual(cfg.chatterbox_speed, 1)

    def test_cross_engine_unknown_fields_rejected(self):
        with self.assertRaises(ValueError):
            Generation(file_ids=['a'], engine={'provider': 'edge', 'cfg_weight': .5})

    def test_quality_requires_mimo(self):
        with self.assertRaises(ServiceError):
            configs_for(Generation(file_ids=[self.file['id']], engine={'provider': 'edge'}, audio_core_enabled=True), self.files, self.root)

    def test_upload_type_size_and_missing_id(self):
        with self.assertRaises(ServiceError):
            self.files.upload('a.py', io.BytesIO(b'code'))
        with self.assertRaises(ServiceError):
            self.files.upload('a.wav', io.BytesIO(b'too big'), limit=1)
        with self.assertRaises(ServiceError):
            self.files.resolve('../etc/passwd')

    def test_registered_file_symlink_replacement_rejected(self):
        path = self.files.resolve(self.file['id'])
        other = self.root / 'other.epub'
        other.write_bytes(b'other')
        path.unlink()
        path.symlink_to(other)
        with self.assertRaises(ServiceError):
            self.files.resolve(self.file['id'])

    def test_confirmation_revision_and_invalidation(self):
        script, _ = _script()
        p = self.projects.put(script)
        p = self.projects.edit(p['id'], ProjectEdit(revision=p['revision'], operation='confirm'))
        self.assertTrue(p['confirmed'])
        self.projects.for_generation(p['id'], p['revision'], 'openai')
        changed = self.projects.edit(p['id'], ProjectEdit(revision=p['revision'], operation='segment',
                      data={'segment_id': 'ch1_s2', 'entity_id': 'narrator'}))
        self.assertFalse(changed['confirmed'])
        with self.assertRaises(ServiceError) as exc:
            self.projects.edit(p['id'], ProjectEdit(revision=p['revision'], operation='confirm'))
        self.assertEqual(exc.exception.status, 409)

    def test_shared_runtime_lock(self):
        first = RuntimeLock(self.root).acquire()
        try:
            with self.assertRaises(ServiceError):
                RuntimeLock(self.root).acquire()
        finally:
            first.close()
        RuntimeLock(self.root).acquire().close()

    def manager(self, runner=fake_worker):
        manager = TaskManager(self.root, self.files, self.projects, runner=runner)
        self.addCleanup(manager.close)
        return manager

    def request(self, request_id='offline-request-1'):
        return TaskRequest(request_id=request_id, generation=Generation(file_ids=[self.file['id']], preview=True))

    def test_spawn_idempotency_result_and_incremental_logs(self):
        manager = self.manager()
        first = manager.start(self.request())
        self.assertEqual(manager.start(self.request())['id'], first['id'])
        deadline = time.time() + 10
        while manager.get(first['id'])['status'] in {'starting', 'running'} and time.time() < deadline:
            time.sleep(.05)
        self.assertEqual(manager.get(first['id'])['status'], 'succeeded')
        self.assertEqual(manager.get(first['id'])['progress']['mode'], 'finished')
        log = manager.logs(first['id'])
        self.assertIn('第二行', log['text'])
        self.assertEqual(manager.logs(first['id'], log['cursor'])['text'], '')
        self.assertEqual(manager.stop(first['id'])['status'], 'succeeded')

    def test_gradio_adapter_mirrors_log_and_uses_same_manager(self):
        manager = self.manager()
        configs = configs_for(self.request().generation, self.files, self.root)
        log = self.root / 'EtA_WebUI_test.log'
        key, process = manager.start_configs(configs, log)
        deadline = time.time() + 10
        while manager.get(key)['status'] in {'starting', 'running'} and time.time() < deadline:
            time.sleep(.05)
        self.assertEqual(manager.get(key)['status'], 'succeeded')
        self.assertIn('第二行', log.read_text())
        self.assertFalse(process.is_alive())

    def test_external_project_edit_invalidates_confirmation(self):
        project = self.projects.put(_script()[0], confirmed=True)
        path = Path(self.projects.entries[project['id']]['path'])
        data = json.loads(path.read_text())
        data['roles']['narrator']['display_name'] = '新旁白'
        atomic_json(path, data)
        changed = self.projects.get(project['id'])
        self.assertFalse(changed['confirmed'])
        self.assertEqual(changed['revision'], project['revision'] + 1)

    def test_canceled_late_result_cannot_overwrite_project(self):
        manager = self.manager(late_result_worker)
        project = self.projects.put(_script()[0])
        key, _ = manager.start_configs(configs_for(self.request().generation, self.files, self.root))
        with manager.lock:
            manager.tasks[key].update(project_id=project['id'], revision=project['revision'])
        self.assertTrue(manager.ready[key].wait(5))
        result = manager.stop(key)
        self.assertEqual(result['status'], 'stopped')
        self.assertEqual(result['progress']['mode'], 'interrupted')
        self.assertEqual(self.projects.get(project['id'])['revision'], project['revision'])

    def test_busy_edit_and_process_group_stop(self):
        manager = self.manager(tree_worker)
        p = self.projects.put(_script()[0])
        task = manager.start(self.request())
        deadline = time.time() + 10
        path = manager.directory / task['id'] / 'EtA_task.log'
        while (not path.exists() or not path.read_text()) and time.time() < deadline:
            time.sleep(.05)
        with self.assertRaises(ServiceError):
            manager.start(self.request('second-request'))
        with self.assertRaises(ServiceError):
            self.projects.edit(p['id'], ProjectEdit(revision=p['revision'], operation='confirm'))
        child = int(path.read_text())
        self.assertEqual(manager.stop(task['id'])['status'], 'stopped')
        self.assertEqual(manager.stop(task['id'])['status'], 'stopped')
        # macOS 不保留僵尸；Linux 的 init 回收可能稍晚。
        deadline = time.time() + 3
        while time.time() < deadline:
            try:
                os.kill(child, 0)
                status = Path(f'/proc/{child}/stat')
                if status.exists() and status.read_text().split()[2] == 'Z':
                    break
                time.sleep(.05)
            except ProcessLookupError:
                break
        else:
            self.fail('停止后仍有存活的子进程')

    def test_restart_marks_interrupted_without_killing_old_pid(self):
        path = self.root / '.studio/tasks/old/task.json'
        atomic_json(path, {'id':'old','request_id':'old-request','status':'running','kind':'generate',
                         'created_at':0,'result':{},'error':''})
        manager = self.manager()
        self.assertEqual(manager.get('old')['status'], 'failed')


class ApiTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.client = TestClient(create_app(self.root, acquire_lock=False))
        self.client.__enter__()
        self.addCleanup(self.client.__exit__, None, None, None)
        self.client.app.state.settings = Settings(self.root)

    def test_bootstrap_no_secrets_and_no_heavy_imports(self):
        value = self.client.get('/api/v1/bootstrap').json()
        self.assertEqual(len(value['providers']), 6)
        self.assertNotIn('admin_password_hash', value['settings'])
        self.assertIsNone(value['current_task'])
        # 全量 suite 可能已为旧 UI 测试导入 Gradio；用独立解释器验证新入口。
        check = subprocess.run([sys.executable, '-c',
            'import sys; from audiobook_generator.api.app import create_app; '
            'assert "audiobook_generator.ui.web_ui" not in sys.modules; '
            'assert "audiobook_generator.tts_providers.voxcpm_tts_provider" not in sys.modules'],
            capture_output=True, text=True, timeout=10)
        self.assertEqual(check.returncode, 0, check.stderr)

    def test_invalid_body_and_unknown_route(self):
        self.assertEqual(self.client.post('/api/v1/tasks', json={}).status_code, 422)
        self.assertEqual(self.client.get('/api/v1/unknown').status_code, 404)
        uploaded = self.client.post('/api/v1/uploads', files={'file': ('bad.docx', b'broken', 'application/octet-stream')}).json()
        parsed = self.client.post('/api/v1/books/chapters', json={'file_id': uploaded['id']})
        self.assertEqual(parsed.status_code, 422)
        self.assertEqual(parsed.json()['error']['code'], 'book_parse_failed')

    def test_cross_origin_mutation_denied(self):
        r = self.client.post('/api/v1/admin/login', headers={'Origin':'https://elsewhere.test'}, json={'password':'admin'})
        self.assertEqual(r.status_code, 403)

    def test_admin_cookie_legacy_upgrade_and_logout(self):
        self.assertEqual(self.client.get('/api/v1/settings/openai').status_code, 401)
        self.assertEqual(self.client.post('/api/v1/admin/login', json={'password':'wrong'}).status_code, 401)
        response = self.client.post('/api/v1/admin/login', json={'password':'admin'})
        self.assertIn('HttpOnly', response.headers['set-cookie'])
        self.assertTrue(self.client.app.state.settings.read()['admin_password_hash'].startswith('pbkdf2$'))
        self.assertEqual(self.client.post('/api/v1/admin/logout').status_code, 200)
        self.assertEqual(self.client.get('/api/v1/settings/openai').status_code, 401)

    def test_upload_range_and_download(self):
        r = self.client.post('/api/v1/uploads', files={'file':('test.wav', b'0123456789', 'audio/wav')})
        self.assertEqual(r.status_code, 200)
        key = r.json()['id']
        response = self.client.get('/api/v1/files/'+key, headers={'Range':'bytes=2-5'})
        self.assertEqual(response.status_code, 206)
        self.assertEqual(response.content, b'2345')
        self.assertIn('attachment', self.client.get('/api/v1/files/'+key+'?download=true').headers['content-disposition'])

    def test_library_delete_ack_and_path_boundary(self):
        (self.root / 'book').mkdir()
        (self.root / 'book/a.wav').write_bytes(b'audio')
        r = self.client.get('/api/v1/library/files', params={'folder':'book'})
        self.assertEqual(len(r.json()), 1)
        self.assertEqual(self.client.get('/api/v1/library/files', params={'folder':'..'}).status_code, 404)
        self.assertEqual(self.client.post('/api/v1/library/delete', json={'folder':'book','names':['a.wav']}).status_code, 422)
        r = self.client.post('/api/v1/library/archive', json={'folder':'book','names':['a.wav']})
        self.assertEqual(r.json()['count'], 1)
        self.assertTrue((self.root/'book/a.wav').exists())

    def test_import_confirm_edit_and_stale_revision(self):
        payload = json.dumps(encode(_script()[0]), ensure_ascii=False).encode()
        file = self.client.post('/api/v1/uploads', files={'file':('role.json',payload)}).json()
        result = self.client.post('/api/v1/projects/import', json={'file_id':file['id']})
        self.assertEqual(result.status_code, 200, result.text)
        p = result.json()
        self.assertTrue(p['confirmed'])
        changed = self.client.patch('/api/v1/projects/'+p['id'], json={'revision':p['revision'],'operation':'lock','data':{'entity_id':'char_1','locked':False}})
        self.assertEqual(changed.status_code, 200)
        self.assertFalse(changed.json()['confirmed'])
        stale = self.client.patch('/api/v1/projects/'+p['id'], json={'revision':p['revision'],'operation':'confirm'})
        self.assertEqual(stale.status_code, 409)

    def test_credentials_file_precedence_empty_preserve_and_mask(self):
        from audiobook_generator.utils import mimo_config
        path = self.root / 'mimo_config.json'
        atomic_json(path, {'api_key':'file-secret','base_url':'https://file.invalid/v1'})
        with patch.object(mimo_config, 'CONFIG_FILE', path), patch.dict(os.environ, {'OPENAI_API_KEY':'env-secret'}):
            value = config_view('openai')
            self.assertTrue(value['configured'])
            self.assertNotIn('api_key', value)
            save_config('openai', {'api_key':'','base_url':'https://changed.invalid/v1'})
            self.assertEqual(json.loads(path.read_text())['api_key'], 'file-secret')
            with self.assertRaises(ServiceError):
                save_config('openai', {'api_key':'••••secret'})


if __name__ == '__main__':
    unittest.main()

class RealParserTaskTest(unittest.TestCase):
    def test_actual_docx_parse_task_without_credentials_or_models(self):
        from docx import Document
        from audiobook_generator.services.books import read_book
        from audiobook_generator.services.models import BookRequest
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / 'sample.docx'
            document = Document()
            document.add_heading('第一章', level=1)
            document.add_paragraph('这是离线迁移验收的正文，确认书籍能够读取并生成目录。')
            document.save(source)
            files = FileStore(root)
            file = files.register(source)
            request = BookRequest(file_id=file['id'])
            catalog, chapters = read_book(request, files)
            self.assertGreater(len(chapters), 0)
            projects = Projects(root, files)
            manager = TaskManager(root, files, projects)
            try:
                value = manager.start(TaskRequest(request_id='actual-parser-test', generation=Generation(
                    file_ids=[file['id']], preview=True, catalog_fingerprint=catalog['fingerprint'])))
                deadline = time.time() + 20
                while manager.get(value['id'])['status'] in {'starting', 'running'} and time.time() < deadline:
                    time.sleep(.05)
                result = manager.get(value['id'])
                self.assertEqual(result['status'], 'succeeded', result['error'])
                self.assertEqual(result['progress']['pct'], 100)
                self.assertIn('预解析完成', manager.logs(value['id'])['text'])
                project = projects.put(_script()[0], confirmed=True)
                with patch('audiobook_generator.services.tasks.require_credentials'), self.assertRaises(ServiceError) as rejected:
                    manager.start(TaskRequest(request_id='wrong-book-project', generation=Generation(
                        file_ids=[file['id']], project_id=project['id'], project_revision=project['revision'])))
                self.assertEqual(rejected.exception.code, 'project_coverage')
                self.assertIsNone(manager.current())
            finally:
                manager.close()
